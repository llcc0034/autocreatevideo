from __future__ import annotations

import hashlib
import io
import re
import zipfile
from pathlib import Path

from django.core.exceptions import ValidationError
from django.core.files.base import ContentFile
from django.db import transaction
from django.db.models import Max
from django.utils import timezone

from studio.models import Asset, ScriptRevision, ScriptSource, SeriesProject
from studio.services.configuration import resolve_configuration
from studio.services.security_boundaries import boundary_int


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
ALLOWED_EXTENSIONS = {
    ".txt": ScriptSource.SourceFormat.PLAIN,
    ".md": ScriptSource.SourceFormat.MARKDOWN,
    ".markdown": ScriptSource.SourceFormat.MARKDOWN,
    ".fountain": ScriptSource.SourceFormat.FOUNTAIN,
    ".docx": ScriptSource.SourceFormat.DOCX,
}
ALLOWED_MIME_BY_FORMAT = {
    ScriptSource.SourceFormat.PLAIN: {"text/plain", "application/octet-stream"},
    ScriptSource.SourceFormat.MARKDOWN: {
        "text/plain", "text/markdown", "text/x-markdown", "application/octet-stream",
    },
    ScriptSource.SourceFormat.FOUNTAIN: {
        "text/plain", "text/fountain", "application/octet-stream",
    },
    ScriptSource.SourceFormat.DOCX: {DOCX_MIME, "application/octet-stream"},
}


class ScriptSourceValidationError(ValidationError):
    def __init__(self, message, *, source: ScriptSource | None = None, code: str = "invalid_script"):
        super().__init__(message, code=code)
        self.source = source
        self.code = code


def _upload_limit(project: SeriesProject) -> int:
    resolved = resolve_configuration(
        project=project,
        required_keys=("script.max_upload_bytes",),
    )
    return int(resolved.values["script.max_upload_bytes"])


def _validate_text(text: str, *, max_bytes: int) -> tuple[str, int, str]:
    normalized = str(text or "").replace("\r\n", "\n").replace("\r", "\n")
    if "\x00" in normalized:
        raise ScriptSourceValidationError("剧本文本包含 NUL 字符", code="script_binary_content")
    if not normalized.strip():
        raise ScriptSourceValidationError("剧本文本为空", code="script_empty")
    payload = normalized.encode("utf-8")
    if len(payload) > max_bytes:
        raise ScriptSourceValidationError(
            f"提取后的剧本文本超过 {max_bytes} 字节上限",
            code="script_too_large",
        )
    return normalized, len(payload), hashlib.sha256(payload).hexdigest()


def _extract_docx(payload: bytes) -> str:
    try:
        with zipfile.ZipFile(io.BytesIO(payload)) as archive:
            members = archive.infolist()
            if len(members) > boundary_int("security.script.docx_max_members"):
                raise ScriptSourceValidationError("DOCX 文件成员过多", code="docx_archive_unsafe")
            total = sum(max(0, item.file_size) for item in members)
            names = {item.filename for item in members}
            if total > boundary_int(
                "security.script.docx_max_uncompressed_bytes",
            ):
                raise ScriptSourceValidationError("DOCX 解压后超过安全上限", code="docx_archive_unsafe")
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise ScriptSourceValidationError("文件不是有效 DOCX", code="docx_invalid")
    except zipfile.BadZipFile as exc:
        raise ScriptSourceValidationError("文件不是有效 DOCX", code="docx_invalid") from exc
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - guarded by requirements
        raise ScriptSourceValidationError(
            "服务器未安装 DOCX 提取依赖",
            code="docx_runtime_unavailable",
        ) from exc
    try:
        document = Document(io.BytesIO(payload))
    except Exception as exc:
        raise ScriptSourceValidationError("DOCX 无法安全解析", code="docx_invalid") from exc
    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(blocks)


def _extract_upload(payload: bytes, source_format: str) -> str:
    if source_format == ScriptSource.SourceFormat.DOCX:
        return _extract_docx(payload)
    try:
        return payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ScriptSourceValidationError(
            "TXT/Markdown/Fountain 必须使用 UTF-8 编码",
            code="script_encoding_invalid",
        ) from exc


def normalize_script_asset_upload(uploaded_file) -> tuple[bytes, dict]:
    """Validate a library script and return canonical UTF-8 Markdown bytes.

    Asset-library scripts are project independent.  They use the same accepted
    formats and safe DOCX extraction as ScriptSource, while the project-level
    size policy and confirmation state remain exclusive to ScriptRevision.
    """

    name = Path(str(getattr(uploaded_file, "name", "") or "")).name
    extension = Path(name).suffix.casefold()
    if extension == ".pdf":
        raise ScriptSourceValidationError(
            "PDF 不属于可复用剧本素材格式；请上传 TXT、Markdown、Fountain 或 DOCX",
            code="pdf_not_supported",
        )
    source_format = ALLOWED_EXTENSIONS.get(extension)
    if source_format is None:
        raise ScriptSourceValidationError(
            "只支持 TXT、Markdown、Fountain 或 DOCX",
            code="script_format_not_supported",
        )
    mime_type = str(
        getattr(uploaded_file, "content_type", "") or "application/octet-stream"
    ).casefold()
    if mime_type not in ALLOWED_MIME_BY_FORMAT[source_format]:
        raise ScriptSourceValidationError(
            "文件扩展名与 MIME 类型不匹配",
            code="script_mime_invalid",
        )
    max_bytes = int(resolve_configuration(
        required_keys=("script.max_upload_bytes",),
    ).values["script.max_upload_bytes"])
    declared_size = int(getattr(uploaded_file, "size", 0) or 0)
    if declared_size > max_bytes:
        raise ScriptSourceValidationError(
            f"剧本文件超过 {max_bytes} 字节上限",
            code="script_too_large",
        )
    uploaded_file.seek(0)
    payload = uploaded_file.read(max_bytes + 1)
    uploaded_file.seek(0)
    if len(payload) > max_bytes:
        raise ScriptSourceValidationError(
            f"剧本文件超过 {max_bytes} 字节上限",
            code="script_too_large",
        )
    if not payload:
        raise ScriptSourceValidationError("上传文件为空", code="script_empty")
    extracted = _extract_upload(payload, source_format)
    normalized, size, digest = _validate_text(extracted, max_bytes=max_bytes)
    canonical = normalized.encode("utf-8")
    return canonical, {
        "manifest_schema_version": 1,
        "source": "user_upload",
        "script_source_format": source_format,
        "original_name": name[:255],
        "original_mime_type": mime_type,
        "original_size_bytes": len(payload),
        "original_sha256": hashlib.sha256(payload).hexdigest(),
        "size_bytes": size,
        "text_sha256": digest,
        "script_normalization": "utf8_markdown_v1",
    }


def _create_revision(*, source, content, created_by, status=ScriptRevision.Status.PREVIEW):
    normalized, size, digest = _validate_text(content, max_bytes=_upload_limit(source.project))
    next_revision = (
        ScriptRevision.objects.filter(source=source).aggregate(value=Max("revision"))["value"]
        or 0
    ) + 1
    return ScriptRevision.objects.create(
        source=source,
        revision=next_revision,
        content=normalized,
        content_size_bytes=size,
        content_sha256=digest,
        status=status,
        created_by=created_by,
    )


@transaction.atomic
def create_pasted_script_source(
    *,
    project: SeriesProject,
    user,
    content: str,
    source_format: str = ScriptSource.SourceFormat.PLAIN,
    source_type: str = ScriptSource.SourceType.PASTE,
) -> tuple[ScriptSource, ScriptRevision]:
    if source_format not in {
        ScriptSource.SourceFormat.PLAIN,
        ScriptSource.SourceFormat.MARKDOWN,
        ScriptSource.SourceFormat.FOUNTAIN,
    }:
        raise ScriptSourceValidationError("粘贴内容只支持 plain/markdown/fountain")
    normalized, size, digest = _validate_text(content, max_bytes=_upload_limit(project))
    source = ScriptSource.objects.create(
        owner=user,
        project=project,
        source_type=source_type,
        source_format=source_format,
        status=ScriptSource.Status.PREVIEW_READY,
        mime_type="text/plain",
        file_size_bytes=size,
        file_sha256=digest,
    )
    revision = _create_revision(source=source, content=normalized, created_by=user)
    return source, revision


def create_uploaded_script_source(
    *,
    project: SeriesProject,
    user,
    uploaded_file,
) -> tuple[ScriptSource, ScriptRevision]:
    name = Path(str(getattr(uploaded_file, "name", "") or "")).name
    extension = Path(name).suffix.casefold()
    if extension == ".pdf":
        raise ScriptSourceValidationError(
            "PDF 不属于本接口支持的可确认剧本格式；请上传 TXT、Markdown、Fountain 或 DOCX",
            code="pdf_not_supported",
        )
    source_format = ALLOWED_EXTENSIONS.get(extension)
    if source_format is None:
        raise ScriptSourceValidationError(
            "只支持 TXT、Markdown、Fountain 或 DOCX",
            code="script_format_not_supported",
        )
    mime_type = str(getattr(uploaded_file, "content_type", "") or "application/octet-stream").casefold()
    if mime_type not in ALLOWED_MIME_BY_FORMAT[source_format]:
        raise ScriptSourceValidationError("文件扩展名与 MIME 类型不匹配", code="script_mime_invalid")
    max_bytes = _upload_limit(project)
    declared_size = int(getattr(uploaded_file, "size", 0) or 0)
    if declared_size > max_bytes:
        raise ScriptSourceValidationError(
            f"剧本文件超过 {max_bytes} 字节上限",
            code="script_too_large",
        )
    payload = uploaded_file.read(max_bytes + 1)
    if len(payload) > max_bytes:
        raise ScriptSourceValidationError(
            f"剧本文件超过 {max_bytes} 字节上限",
            code="script_too_large",
        )
    if not payload:
        raise ScriptSourceValidationError("上传文件为空", code="script_empty")
    source = ScriptSource.objects.create(
        owner=user,
        project=project,
        source_type=ScriptSource.SourceType.UPLOAD,
        source_format=source_format,
        status=ScriptSource.Status.PENDING,
        original_name=name[:255],
        mime_type=mime_type,
        file_size_bytes=len(payload),
        file_sha256=hashlib.sha256(payload).hexdigest(),
    )
    source.file.save(name, ContentFile(payload), save=False)
    source.save(update_fields=("file", "updated_at"))
    try:
        extracted = _extract_upload(payload, source_format)
        revision = _create_revision(source=source, content=extracted, created_by=user)
    except ScriptSourceValidationError as exc:
        source.status = ScriptSource.Status.FAILED
        source.validation_message = str(exc)[:2000]
        source.save(update_fields=("status", "validation_message", "updated_at"))
        exc.source = source
        raise
    source.status = ScriptSource.Status.PREVIEW_READY
    source.validation_message = "文本提取完成，等待确认。"
    source.save(update_fields=("status", "validation_message", "updated_at"))
    return source, revision


@transaction.atomic
def create_script_source_from_asset(
    *,
    project: SeriesProject,
    user,
    asset: Asset,
) -> tuple[ScriptSource, ScriptRevision]:
    """Import a reusable script asset as a project preview revision."""

    asset = Asset.objects.select_for_update().get(pk=asset.pk)
    if not asset.is_active or asset.kind != Asset.Kind.SCRIPT or not asset.file:
        raise ScriptSourceValidationError(
            "所选素材不是可用的剧本文件",
            code="script_asset_unavailable",
        )
    if asset.owner_id not in (None, user.id) and not user.is_staff:
        raise ScriptSourceValidationError(
            "不能导入其他用户的剧本素材",
            code="script_asset_forbidden",
        )
    max_bytes = _upload_limit(project)
    try:
        if asset.file.size > max_bytes:
            raise ScriptSourceValidationError(
                f"剧本素材超过项目 {max_bytes} 字节上限",
                code="script_too_large",
            )
        with asset.file.open("rb") as handle:
            payload = handle.read(max_bytes + 1)
    except ScriptSourceValidationError:
        raise
    except (OSError, ValueError) as exc:
        raise ScriptSourceValidationError(
            "剧本素材文件不可读取",
            code="script_asset_unavailable",
        ) from exc
    if len(payload) > max_bytes:
        raise ScriptSourceValidationError(
            f"剧本素材超过项目 {max_bytes} 字节上限",
            code="script_too_large",
        )
    payload_sha256 = hashlib.sha256(payload).hexdigest()
    if (
        asset.mime_type != "text/markdown"
        or not re.fullmatch(r"[0-9a-f]{64}", str(asset.checksum or ""))
        or payload_sha256 != asset.checksum
    ):
        raise ScriptSourceValidationError(
            "剧本素材字节与受管完整性记录不一致，拒绝导入",
            code="script_asset_integrity_mismatch",
        )
    try:
        content = payload.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ScriptSourceValidationError(
            "剧本素材不是规范 UTF-8 文本",
            code="script_encoding_invalid",
        ) from exc
    normalized, size, digest = _validate_text(content, max_bytes=max_bytes)
    source = ScriptSource.objects.create(
        owner=user,
        project=project,
        source_asset=asset,
        source_type=ScriptSource.SourceType.LIBRARY_ASSET,
        source_format=ScriptSource.SourceFormat.MARKDOWN,
        status=ScriptSource.Status.PREVIEW_READY,
        original_name=Path(asset.file.name).name[:255],
        mime_type=asset.mime_type or "text/markdown",
        file_size_bytes=size,
        file_sha256=digest,
        validation_message="已从素材库导入，等待人工确认。",
    )
    revision = _create_revision(source=source, content=normalized, created_by=user)
    return source, revision


@transaction.atomic
def create_script_revision(*, source: ScriptSource, user, content: str) -> ScriptRevision:
    source = ScriptSource.objects.select_for_update().get(pk=source.pk)
    if source.status == ScriptSource.Status.FAILED:
        raise ScriptSourceValidationError("失败的 source 不能直接新增 revision")
    if source.source_type == ScriptSource.SourceType.AI_GENERATED:
        raise ScriptSourceValidationError(
            "AI 生成的 source 只能通过对应 ProductionRun 审核，不能从公共 revision API 派生",
            code="production_run_review_required",
        )
    revision = _create_revision(source=source, content=content, created_by=user)
    source.status = ScriptSource.Status.PREVIEW_READY
    source.validation_message = "已创建新的文本预览 revision。"
    source.save(update_fields=("status", "validation_message", "updated_at"))
    return revision


@transaction.atomic
def confirm_script_revision(
    revision: ScriptRevision,
    *,
    user,
    allow_generated: bool = False,
) -> ScriptRevision:
    revision = ScriptRevision.objects.select_for_update().select_related("source").get(pk=revision.pk)
    source = ScriptSource.objects.select_for_update().get(pk=revision.source_id)
    if (
        revision.source.source_type == ScriptSource.SourceType.AI_GENERATED
        or revision.origin_generation_attempt_id
    ) and not allow_generated:
        raise ScriptSourceValidationError(
            "AI 生成的 preview 必须通过 ProductionRun script_review 审批确认",
            code="production_run_review_required",
        )
    if revision.status != ScriptRevision.Status.PREVIEW:
        raise ScriptSourceValidationError("只有 preview revision 可以确认")
    ScriptRevision.objects.filter(
        source=source,
        status=ScriptRevision.Status.CONFIRMED,
    ).exclude(pk=revision.pk).update(status=ScriptRevision.Status.SUPERSEDED, updated_at=timezone.now())
    revision.status = ScriptRevision.Status.CONFIRMED
    revision.confirmed_by = user
    revision.confirmed_at = timezone.now()
    revision.save(update_fields=("status", "confirmed_by", "confirmed_at", "updated_at"))
    source.status = ScriptSource.Status.CONFIRMED
    source.validation_message = "文本 revision 已确认，可进入 ProductionRun。"
    source.save(update_fields=("status", "validation_message", "updated_at"))
    return revision
