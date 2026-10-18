from io import BytesIO
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.files.base import ContentFile
from django.test import TestCase, override_settings
from rest_framework.test import APIClient

from studio.models import (
    Asset,
    CharacterBible,
    ConfigDefinition,
    ConfigRevision,
    Episode,
    GeneratedArtifact,
    PropBible,
    ProductionRecipeVersion,
    ProductionRun,
    ScriptRevision,
    ScriptSource,
    SceneBible,
    SeriesProject,
    Shot,
    StageRun,
)
from studio.services.json_contracts import assign_json_contract
from studio.services.configuration import resolve_configuration
from studio.services.production_runs import (
    advance_production_run,
    StageBlocked,
    _assert_frozen_formal_objects,
    _canonical_storyboard,
    _sha256,
)
from studio.tests import test_production_recipe_contracts as route_fixtures
from studio.tests.structured_bible_fixtures import confirmed_structured_bible_source


@override_settings(AI_PROVIDER_ALLOWED_HOSTS=("api.example.com", "cdn.example.com"))
class ProductionRunApiTests(TestCase):
    setUpTestData = classmethod(route_fixtures.RecipeExecutionRouteTests.setUpTestData.__func__)
    write_contract = staticmethod(route_fixtures.RecipeExecutionRouteTests.write_contract)

    def setUp(self):
        route_fixtures.RecipeExecutionRouteTests.setUp(self)
        self.user = get_user_model().objects.create_user(
            username="production-owner",
            password="StrongPass123!",
        )
        self.project = SeriesProject.objects.create(language="Chinese", aspect_ratio="9:16",
            owner=self.user,
            title="一键生产测试",
            width=768, height=1344, fps=24, default_shot_duration_ms=5000,
            continuity_seed=1,
            production_settings={"default_conditioning_mode": "T2VA"},
        )
        self.recipe = ProductionRecipeVersion(
            recipe_key="test-script-review",
            version=1,
            name="Test script review",
            status=ProductionRecipeVersion.Status.PUBLISHED,
        )
        assign_json_contract(self.recipe, "definition", self._recipe_definition([
                {"key": "script_input", "mode": "confirmed_revision", "approval": "auto"},
                {"key": "bible_review", "mode": "manual", "approval": "required"},
                {"key": "storyboard", "mode": "deterministic", "approval": "auto"},
                {"key": "storyboard_review", "mode": "review", "approval": "required"},
            ]), schema_version=3)
        self.recipe.save()
        self.client = APIClient()
        self.client.force_authenticate(self.user)

    def _recipe_definition(self, stages):
        return {"schema_version": 3, "interaction_mode": "assisted", "stages": [*stages,
            {"key": "media_generation", "mode": "existing_pipeline", "approval": "auto",
                "routing": {"mode": "fixed", "routes": [deepcopy(self.route)], "variant_count": 1, "evaluator": None}},
            {"key": "artifact_selection", "mode": "review", "approval": "required"},
            {"key": "assembly", "mode": "existing_pipeline", "approval": "auto"},
            {"key": "delivery", "mode": "existing_pipeline", "approval": "auto"}]}

    def _confirmed_script(self):
        # Original M2 structure writer and confirmation, not a canonical-text
        # shortcut or fabricated succeeded Stage.
        revision = confirmed_structured_bible_source(
            SimpleNamespace(actor=self.user, project=self.project), entities=[],
            episodes=[{"episode": 1, "title": "第一集", "shots": [
                {"shot": 1, "duration_ms": 6000, "setting": "雨夜老宅", "action": "沈砚推门而入。"}
            ]}])
        self.assertEqual(revision.status, ScriptRevision.Status.CONFIRMED)
        self.assertEqual(revision.source.source_type, ScriptSource.SourceType.PASTE)
        return str(revision.pk)


    def _create_run(self, revision_id, *, key="run-1", episode_count=None, episode_ids=None):
        payload = {
            "project_id": str(self.project.id),
            "script_revision_id": revision_id,
            "plan_source": "recipe",
            "recipe_version_id": str(self.recipe.id),
            "idempotency_key": key,
        }
        if episode_count is not None:
            payload["episode_count"] = episode_count
        if episode_ids is not None:
            payload["episode_ids"] = episode_ids
        response = self.client.post("/api/v1/production-runs/", payload, format="json")
        if response.status_code == 201:
            # Simulate the normal platform worker's next tick. The returned
            # API receipt is deliberately not rewritten to its later state.
            advance_production_run(ProductionRun.objects.get(pk=response.data["id"]))
        return response

    def test_imported_script_keeps_generation_counts_unset_not_project_defaults(self):
        revision_id = self._confirmed_script()

        created = self._create_run(revision_id, key="creation-policy-episode-count")

        self.assertEqual(created.status_code, 201, created.data)
        run = ProductionRun.objects.get(pk=created.data["id"])
        inputs = run.input_snapshot["request_intent"]["creation_inputs"]
        self.assertIsNone(inputs["episode_count"])
        self.assertIsNone(inputs["shots_per_episode"])
        self.assertIsNone(inputs["episode_duration_seconds"])

    def test_partial_generation_counts_are_rejected_without_creating_a_run(self):
        revision_id = self._confirmed_script()
        response = self._create_run(revision_id, episode_count=2, key="partial-counts")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertIn("完整提供", str(response.data))
        self.assertFalse(ProductionRun.objects.filter(idempotency_key="partial-counts").exists())


    def test_public_snapshot_projections_resolve_redaction_boundary_once(self):
        from studio import platform_serializers
        from studio.services import configuration as configuration_service

        snapshot = {
            "schema_version": 1,
            "values": {
                "safe": {
                    "usage": {"prompt_tokens": 6, "completion_tokens": 4},
                },
            },
            "sources": {
                f"key-{index}": {
                    "source": "system",
                    "scope": "system",
                    "source_type": "system",
                    "version": 1,
                    "revision": 1,
                }
                for index in range(500)
            },
        }
        routing = {
            "schema_version": 1,
            "mode": "fixed",
            "variant_count": 1,
            "routes": [
                {
                    "index": index,
                    "model_identifier": f"model-{index}",
                    "driver_type": "comfyui",
                    "is_external": False,
                }
                for index in range(100)
            ],
        }
        original = configuration_service.resolve_security_boundaries

        with patch.object(
            configuration_service,
            "resolve_security_boundaries",
            wraps=original,
        ) as resolver:
            public_snapshot = platform_serializers.public_configuration_snapshot(
                snapshot,
            )
        self.assertEqual(resolver.call_count, 1)
        self.assertEqual(len(public_snapshot["sources"]), 500)

        with patch.object(
            configuration_service,
            "resolve_security_boundaries",
            wraps=original,
        ) as resolver:
            public_routing = platform_serializers.public_stage_routing(routing)
        self.assertEqual(resolver.call_count, 1)
        self.assertEqual(len(public_routing["routes"]), 100)


    def test_stage_snapshots_larger_than_default_json_limit_remain_serializable(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="large-public-stage-snapshot")
        self.assertEqual(created.status_code, 201, created.data)
        stage = StageRun.objects.get(
            production_run_id=created.data["id"],
            stage_key="bible_review",
        )
        assign_json_contract(
            stage,
            "output_snapshot",
            {"schema_version": 1, "large_audit_payload": "x" * 70_000},
            max_bytes=8 * 1024 * 1024,
        )
        stage.save(update_fields=(
            "output_snapshot",
            "output_snapshot_schema_version",
            "output_snapshot_size_bytes",
            "output_snapshot_sha256",
            "updated_at",
        ))

        detail = self.client.get(
            f"/api/v1/production-runs/{created.data['id']}/",
        )

        self.assertEqual(detail.status_code, 200, detail.data)
        serialized = next(
            item for item in detail.data["stages"]
            if item["stage_key"] == "bible_review"
        )
        self.assertEqual(
            serialized["output_snapshot"]["large_audit_payload"],
            "x" * 70_000,
        )
        self.assertEqual(len(serialized["output_snapshot_sha256"]), 64)

    def test_confirmed_script_enters_real_gate_and_idempotency_is_enforced(self):
        revision_id = self._confirmed_script()

        created = self._create_run(revision_id)

        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(created.data["status"], ProductionRun.Status.QUEUED)
        current = self.client.get(f"/api/v1/production-runs/{created.data['id']}/")
        self.assertEqual(current.data["status"], ProductionRun.Status.AWAITING_APPROVAL,
            (current.data["blocked_code"], current.data["blocked_reason"]))
        self.assertEqual(current.data["current_stage_key"], "bible_review")
        self.assertEqual(current.data["blocked_code"], "manual_stage_confirmation_required")
        self.assertEqual(created.data["config_snapshot"]["schema_version"], 1)
        run_id = created.data["id"]
        artifact = GeneratedArtifact.objects.get(
            production_run_id=run_id,
            kind=GeneratedArtifact.Kind.SCRIPT,
        )
        self.assertEqual(artifact.status, GeneratedArtifact.Status.SELECTED)
        self.assertEqual(len(artifact.manifest_metadata_sha256), 64)

        replayed = self._create_run(revision_id)
        self.assertEqual(replayed.status_code, 200, replayed.data)
        self.assertEqual(replayed.data["id"], run_id)
        conflict = self._create_run(revision_id, episode_count=2)
        self.assertEqual(conflict.status_code, 409, conflict.data)
        self.assertEqual(
            conflict.data["error"]["code"],
            "production_run_conflict",
        )
        self.assertEqual(ProductionRun.objects.filter(owner=self.user).count(), 1)

        unknown = self.client.post(f"/api/v1/production-runs/{run_id}/approve/", {
            "stage_key": "not-a-stage",
            "decision": "approve",
            "comment": "invalid",
        }, format="json")
        self.assertEqual(unknown.status_code, 400, unknown.data)

        detail = self.client.get(f"/api/v1/production-runs/{run_id}/bible-review/")
        self.assertEqual(detail.status_code, 200, detail.data)
        storyboard = self.client.post(f"/api/v1/production-runs/{run_id}/approve/", {
            "stage_key": "bible_review",
            "decision": "approve",
            "bible_review": {"expected_source_snapshot_sha256": detail.data["source_snapshot_sha256"],
                "expected_project_bible_sha256": detail.data["project_bible_sha256"], "reviewed_candidate_ids": [],
                "reviewed_setting_keys": [], "reviewed_project_bible": True},
            "comment": "本轮无实体参考，已明确核对空 Bible",
        }, format="json")
        # This exact API target does not require entity-reference media, and
        # the confirmed script declares no entities. Do not revive the old
        # unconditional requirement for all three Bible categories.
        self.assertEqual(storyboard.status_code, 200, storyboard.data.get("error"))
        for model in (CharacterBible, SceneBible, PropBible):
            self.assertFalse(model.objects.filter(project=self.project).exists())
        self.assertEqual(storyboard.data["current_stage_key"], "storyboard_review")
        self.assertEqual(storyboard.data["status"], ProductionRun.Status.AWAITING_APPROVAL)

        source_snapshot_sha256 = next(
            item["output_snapshot_sha256"]
            for item in storyboard.data["stages"]
            if item["stage_key"] == "storyboard"
        )
        selection = self.client.get(f"/api/v1/production-runs/{run_id}/storyboard-reference-selection/")
        self.assertEqual(selection.status_code, 200, selection.data)
        self.assertEqual(selection.data["source_snapshot_sha256"], source_snapshot_sha256)
        approved_storyboard = self.client.post(
            f"/api/v1/production-runs/{run_id}/approve/",
            {
                "stage_key": "storyboard_review",
                "decision": "approve",
                "comment": "分镜冻结证据已核对",
                "expected_source_snapshot_sha256": source_snapshot_sha256,
                "expected_selection_sha256": selection.data["selection_sha256"],
            },
            format="json",
        )
        self.assertEqual(approved_storyboard.status_code, 200, approved_storyboard.data)
        self.assertEqual(
            ProductionRun.objects.get(pk=run_id).stages.get(
                stage_key="storyboard_review",
            ).approval_decisions.count(),
            1,
        )

    def test_library_script_preview_is_rejected_until_explicit_confirmation(self):
        payload = "# 第一集\n\n素材库剧本必须先人工确认。".encode("utf-8")
        asset = Asset(
            owner=self.user,
            kind=Asset.Kind.SCRIPT,
            name="素材库门禁剧本",
            mime_type="text/markdown",
            checksum=hashlib.sha256(payload).hexdigest(),
            metadata={"script_normalization": "utf8_markdown_v1"},
        )
        asset.file.save("library-gate.md", ContentFile(payload), save=True)
        stored_path = asset.file.path
        self.addCleanup(lambda: Path(stored_path).unlink(missing_ok=True))
        imported = self.client.post(
            "/api/v1/script-sources/",
            {"project_id": str(self.project.id), "asset_id": str(asset.id)},
            format="json",
        )
        self.assertEqual(imported.status_code, 201, imported.data)
        revision_id = imported.data["revision"]["id"]
        self.assertEqual(imported.data["revision"]["status"], ScriptRevision.Status.PREVIEW)
        rejected = self._create_run(revision_id, key="library-preview-rejected")
        self.assertEqual(rejected.status_code, 400, rejected.data)
        confirmed = self.client.post(
            f"/api/v1/script-revisions/{revision_id}/confirm/",
            {},
            format="json",
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.data)
        accepted = self._create_run(revision_id, key="library-preview-confirmed")
        self.assertEqual(accepted.status_code, 201, accepted.data)
        self.assertEqual(
            str(ProductionRun.objects.get(pk=accepted.data["id"]).script_revision_id),
            str(revision_id),
        )

    def test_existing_episode_and_automatic_storyboard_are_rejected_before_run_creation(self):
        revision_id = self._confirmed_script()
        Episode.objects.create(prompt_settings={"include_audio": True}, project=self.project, number=1, title="既有第一集")
        run_revision_count = ConfigRevision.objects.filter(
            scope_type=ConfigRevision.ScopeType.RUN,
        ).count()

        rejected = self._create_run(
            revision_id,
            key="existing-episode-auto-storyboard",
        )

        self.assertEqual(rejected.status_code, 400, rejected.data)
        self.assertEqual(
            rejected.data["error"]["code"],
            "existing_storyboard_requires_manual_resolution",
        )
        self.assertIn("storyboard/manual", str(rejected.data))
        self.assertFalse(
            ProductionRun.objects.filter(
                idempotency_key="existing-episode-auto-storyboard",
            ).exists(),
        )
        self.assertEqual(
            ConfigRevision.objects.filter(
                scope_type=ConfigRevision.ScopeType.RUN,
            ).count(),
            run_revision_count,
        )

    def test_existing_episode_requires_manual_storyboard_and_rejects_skip(self):
        revision_id = self._confirmed_script()
        Episode.objects.create(prompt_settings={"include_audio": True}, project=self.project, number=1, title="既有第一集")

        for index, mode in enumerate(("manual", "skip"), start=1):
            with self.subTest(mode=mode):
                recipe = ProductionRecipeVersion(
                    recipe_key=f"existing-storyboard-{mode}",
                    version=1,
                    name=f"Existing storyboard {mode}",
                    status=ProductionRecipeVersion.Status.PUBLISHED,
                )
                assign_json_contract(recipe, "definition", self._recipe_definition([
                        {
                            "key": "script_input",
                            "mode": "confirmed_revision",
                            "approval": "auto",
                        },
                        {
                            "key": "storyboard",
                            "mode": mode,
                            "approval": "required" if mode == "manual" else "auto",
                        },
                    ]), schema_version=3)
                recipe.save()
                response = self.client.post("/api/v1/production-runs/", {
                    "project_id": str(self.project.id),
                    "script_revision_id": revision_id,
                    "plan_source": "recipe",
                    "recipe_version_id": str(recipe.id),
                    "idempotency_key": f"existing-storyboard-{index}",
                }, format="json")

                self.assertEqual(response.status_code, 201 if mode == "manual" else 400, response.data)
                self.assertEqual(ProductionRun.objects.filter(idempotency_key=f"existing-storyboard-{index}").exists(), mode == "manual")

    def test_downstream_manual_input_freezes_selected_existing_episodes_without_script(self):
        selected_episode = Episode.objects.create(prompt_settings={"include_audio": True},
            project=self.project,
            number=1,
            title="沿用第一集",
            status=Episode.Status.DRAFT,
        )
        selected_shot = Shot.objects.create(include_audio=True,
            episode=selected_episode,
            number=1,
            title="既有镜头",

            duration_ms=self.project.default_shot_duration_ms,
            media_type=Shot.MediaType.VIDEO,
            conditioning_mode="T2VA",
        )
        Episode.objects.create(prompt_settings={"include_audio": True},
            project=self.project,
            number=2,
            title="未选择且尚无分镜的第二集",
            status=Episode.Status.DRAFT,
        )
        recipe = ProductionRecipeVersion(
            recipe_key="existing-episode-snapshot",
            version=1,
            name="Existing episode snapshot",
            status=ProductionRecipeVersion.Status.PUBLISHED,
        )
        assign_json_contract(recipe, "definition", self._recipe_definition([
                {"key": "script_input", "mode": "manual", "approval": "required"},
                {"key": "storyboard", "mode": "manual", "approval": "required"},
            ]), schema_version=3)
        recipe.save()
        created = self.client.post("/api/v1/production-runs/", {
            "project_id": str(self.project.id),
            "plan_source": "recipe",
            "recipe_version_id": str(recipe.id),
            "episode_ids": [str(selected_episode.id)],
            "idempotency_key": "existing-episode-snapshot",
        }, format="json")
        self.assertEqual(created.status_code, 201, created.data)
        run = advance_production_run(ProductionRun.objects.get(pk=created.data["id"]))
        self.assertEqual(run.current_stage_key, "script_input")
        self.assertEqual(run.status, ProductionRun.Status.AWAITING_APPROVAL)

        script_gate = self.client.post(
            f"/api/v1/production-runs/{created.data['id']}/approve/",
            {
                "stage_key": "script_input",
                "decision": "approve",
                "comment": "已核对现有分集范围",
            },
            format="json",
        )

        self.assertEqual(script_gate.status_code, 200, script_gate.data)
        self.assertEqual(script_gate.data["current_stage_key"], "storyboard")
        run = ProductionRun.objects.get(pk=created.data["id"])
        self.assertIsNone(run.script_revision_id)
        script_stage = run.stages.get(stage_key="script_input")
        self.assertEqual(
            script_stage.output_snapshot["formal_object_type"],
            "existing_episode_selection",
        )
        self.assertEqual(
            script_stage.output_snapshot["canonical"]["episodes"],
            [{
                "id": str(selected_episode.id),
                "number": selected_episode.number,
            }],
        )

        storyboard_gate = self.client.post(
            f"/api/v1/production-runs/{run.id}/approve/",
            {
                "stage_key": "storyboard",
                "decision": "approve",
                "comment": "已核对选中分集的正式镜头",
            },
            format="json",
        )
        self.assertEqual(storyboard_gate.status_code, 200, storyboard_gate.data)
        storyboard_stage = run.stages.get(stage_key="storyboard")
        self.assertEqual(storyboard_stage.status, StageRun.Status.SUCCEEDED)
        self.assertEqual(
            storyboard_stage.output_snapshot["formal_object_type"],
            "storyboard",
        )
        frozen_shots = storyboard_stage.output_snapshot["canonical"]["episodes"][0]["shots"]
        self.assertEqual([row["id"] for row in frozen_shots], [str(selected_shot.id)])

    def test_retrieve_run_is_read_only_and_never_advances_the_orchestrator(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="read-only-run-detail")
        self.assertEqual(created.status_code, 201, created.data)
        run = ProductionRun.objects.get(pk=created.data["id"])
        run_fact = (run.status, run.current_stage_key, run.updated_at)
        stage_facts = list(
            run.stages.order_by("ordinal").values_list(
                "id", "status", "updated_at",
            )
        )

        with patch("studio.platform_views.advance_production_run") as advance:
            detail = self.client.get(f"/api/v1/production-runs/{run.id}/")

        self.assertEqual(detail.status_code, 200, detail.data)
        advance.assert_not_called()
        run.refresh_from_db()
        self.assertEqual((run.status, run.current_stage_key, run.updated_at), run_fact)
        self.assertEqual(
            list(
                run.stages.order_by("ordinal").values_list(
                    "id", "status", "updated_at",
                )
            ),
            stage_facts,
        )

    def test_storyboard_review_rejects_creative_drift_without_writing_decision(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="storyboard-approval-drift")
        self.assertEqual(created.status_code, 201, created.data)
        run = ProductionRun.objects.get(pk=created.data["id"])
        bible = self.client.get(f"/api/v1/production-runs/{run.id}/bible-review/")
        self.assertEqual(bible.status_code, 200, bible.data)
        advanced = self.client.post(f"/api/v1/production-runs/{run.id}/approve/", {
            "stage_key": "bible_review", "decision": "approve",
            "bible_review": {"expected_source_snapshot_sha256": bible.data["source_snapshot_sha256"],
                "expected_project_bible_sha256": bible.data["project_bible_sha256"],
                "reviewed_candidate_ids": [], "reviewed_setting_keys": [], "reviewed_project_bible": True},
        }, format="json")
        self.assertEqual(advanced.status_code, 200, advanced.data)
        self.assertEqual(advanced.data["current_stage_key"], "storyboard_review")
        source_stage = run.stages.get(stage_key="storyboard")
        review_stage = run.stages.get(stage_key="storyboard_review")
        source_contract_hash = source_stage.output_snapshot_sha256
        # Use the real writer and stage-input receipts, then submit precisely
        # what was observed before the authored content changed.
        selection = self.client.get(f"/api/v1/production-runs/{run.id}/storyboard-reference-selection/")
        self.assertEqual(selection.status_code, 200, selection.data)
        shot = Shot.objects.get(episode__project=self.project)
        shot.title = "Changed after the source and selection were reviewed"
        shot.save(update_fields=("title", "updated_at"))
        rejected = self.client.post(f"/api/v1/production-runs/{run.id}/approve/", {
            "stage_key": "storyboard_review", "decision": "approve",
            "comment": "不应静默重新冻结",
            "expected_source_snapshot_sha256": selection.data["source_snapshot_sha256"],
            "expected_selection_sha256": selection.data["selection_sha256"],
        }, format="json")

        self.assertEqual(rejected.status_code, 409, rejected.data)
        self.assertEqual(rejected.data["error"]["code"], "storyboard_snapshot_drift")
        review_stage.refresh_from_db()
        source_stage.refresh_from_db()
        self.assertEqual(review_stage.status, StageRun.Status.AWAITING_APPROVAL)
        self.assertEqual(review_stage.approval_decisions.count(), 0)
        self.assertEqual(source_stage.output_snapshot_sha256, source_contract_hash)

    def test_reject_and_cancel_are_durable_terminal_transitions(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="reject-run")
        run_id = created.data["id"]
        rejected = self.client.post(f"/api/v1/production-runs/{run_id}/approve/", {
            "stage_key": "bible_review",
            "decision": "reject",
            "comment": "需要重写",
        }, format="json")
        self.assertEqual(rejected.status_code, 200, rejected.data)
        self.assertEqual(rejected.data["status"], ProductionRun.Status.FAILED)

        second = self._create_run(revision_id, key="cancel-run")
        cancelled = self.client.post(
            f"/api/v1/production-runs/{second.data['id']}/cancel/",
            {},
            format="json",
        )
        self.assertEqual(cancelled.status_code, 200, cancelled.data)
        self.assertEqual(cancelled.data["status"], ProductionRun.Status.CANCELLED)

    def test_run_override_rejects_secret_even_when_definition_claims_run_scope(self):
        definition = ConfigDefinition(
            key="provider.test.secret",
            domain="provider",
            value_type=ConfigDefinition.ValueType.STRING,
            allowed_scopes="run",
            is_secret=True,
        )
        assign_json_contract(definition, "value_schema", {})
        definition.save()
        revision_id = self._confirmed_script()
        response = self.client.post("/api/v1/production-runs/", {
            "project_id": str(self.project.id),
            "script_revision_id": revision_id,
            "plan_source": "recipe",
            "recipe_version_id": str(self.recipe.id),
            "idempotency_key": "secret-run",
            "config_overrides": {
                definition.key: {"credential_id": "11111111-1111-4111-8111-111111111111"},
            },
        }, format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertNotIn("credential_id", json.dumps(response.data, default=str))
        self.assertNotIn("11111111-1111-4111-8111-111111111111", json.dumps(response.data, default=str))
        self.assertFalse(ProductionRun.objects.filter(idempotency_key="secret-run").exists())

    def test_run_overrides_drive_fields_and_frozen_effective_snapshot(self):
        revision_id = self._confirmed_script()
        response = self.client.post("/api/v1/production-runs/", {
            "project_id": str(self.project.id),
            "script_revision_id": revision_id,
            "plan_source": "recipe",
            "recipe_version_id": str(self.recipe.id),
            "idempotency_key": "effective-run-overrides",
            "config_overrides": {
                "production.allow_external_calls": False,
            },
        }, format="json")
        self.assertEqual(response.status_code, 201, response.data)
        run = ProductionRun.objects.get(pk=response.data["id"])
        from studio.services.execution_resource_policy import RESOURCE_POLICY, read_resource_policy
        self.assertNotIn("max_generation_tasks", {field.name for field in ProductionRun._meta.fields})
        self.assertFalse(run.allow_external_calls)
        self.assertNotIn("production.max_generation_tasks", run.config_snapshot["values"])
        self.assertFalse(run.config_snapshot["values"]["production.allow_external_calls"])
        self.assertEqual(
            run.config_snapshot["sources"][RESOURCE_POLICY]["scope"],
            ConfigRevision.ScopeType.SYSTEM,
        )

    def test_invalid_nested_sensitive_override_is_prevalidated_without_rows(self):
        definition = ConfigDefinition(
            key="production.test.payload",
            domain="production",
            value_type=ConfigDefinition.ValueType.OBJECT,
            allowed_scopes="run",
        )
        assign_json_contract(definition, "value_schema", {"type": "object"})
        definition.save()
        revision_id = self._confirmed_script()
        response = self.client.post("/api/v1/production-runs/", {
            "project_id": str(self.project.id),
            "script_revision_id": revision_id,
            "plan_source": "recipe",
            "recipe_version_id": str(self.recipe.id),
            "idempotency_key": "nested-secret-run",
            "config_overrides": {
                definition.key: {"access_token": "must-not-persist"},
            },
        }, format="json")
        self.assertEqual(response.status_code, 400, response.data)
        self.assertNotIn("must-not-persist", json.dumps(response.data, default=str))
        self.assertFalse(ProductionRun.objects.filter(idempotency_key="nested-secret-run").exists())
        self.assertFalse(ConfigRevision.objects.filter(scope_type="run").exists())

    def test_stage_initialization_failure_rolls_back_run_revision_and_stages(self):
        revision_id = self._confirmed_script()
        with patch.object(StageRun.objects, "bulk_create", side_effect=RuntimeError("stage failure")):
            with self.assertRaisesRegex(RuntimeError, "stage failure"):
                self.client.post("/api/v1/production-runs/", {
                    "project_id": str(self.project.id),
                    "script_revision_id": revision_id,
                    "plan_source": "recipe",
                    "recipe_version_id": str(self.recipe.id),
                    "idempotency_key": "stage-rollback-run",
                    "config_overrides": {"production.allow_external_calls": False},
                }, format="json")
        self.assertFalse(ProductionRun.objects.filter(idempotency_key="stage-rollback-run").exists())
        self.assertFalse(ConfigRevision.objects.filter(scope_type="run").exists())
        self.assertFalse(StageRun.objects.filter(production_run__idempotency_key="stage-rollback-run").exists())

    def test_idempotent_replay_rejects_incomplete_initialization(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="incomplete-idempotent-run")
        self.assertEqual(created.status_code, 201, created.data)
        run = ProductionRun.objects.get(pk=created.data["id"])
        run.stages.order_by("-ordinal").first().delete()

        replay = self._create_run(revision_id, key="incomplete-idempotent-run")
        self.assertEqual(replay.status_code, 409, replay.data)
        self.assertEqual(
            replay.data["error"]["code"],
            "production_run_conflict",
        )
        self.assertIn("冻结计划不完整", str(replay.data))

    def test_episode_selection_is_owned_deduplicated_sorted_and_frozen(self):
        first = Episode.objects.create(prompt_settings={"include_audio": True}, project=self.project, number=2, title="第二集")
        second = Episode.objects.create(prompt_settings={"include_audio": True}, project=self.project, number=1, title="第一集")
        definition = json.loads(json.dumps(self.recipe.definition))
        storyboard = next(
            item for item in definition["stages"] if item["key"] == "storyboard"
        )
        storyboard["mode"] = "manual"
        storyboard["approval"] = "required"
        definition["stages"][0].update(mode="manual", approval="required")
        assign_json_contract(self.recipe, "definition", definition, schema_version=3)
        self.recipe.save(update_fields=(
            "definition", "definition_schema_version", "definition_size_bytes",
            "definition_sha256", "updated_at",
        ))
        foreign_owner = get_user_model().objects.create_user(
            username="foreign-episode-owner", password="StrongPass123!",
        )
        foreign_project = SeriesProject.objects.create(language="Chinese", aspect_ratio="9:16", owner=foreign_owner, title="其他项目", width=768, height=1344, fps=24, default_shot_duration_ms=5000, continuity_seed=1)
        foreign = Episode.objects.create(prompt_settings={"include_audio": True}, project=foreign_project, number=1, title="外部集")
        revision_id = self._confirmed_script()

        invalid = self._create_run(
            None,
            key="foreign-episode-run",
            episode_ids=[str(first.id), str(foreign.id)],
        )
        self.assertEqual(invalid.status_code, 400, invalid.data)
        created = self._create_run(
            None,
            key="selected-episodes-run",
            episode_ids=[str(first.id), str(second.id), str(first.id)],
        )
        self.assertEqual(created.status_code, 201, created.data)
        self.assertEqual(
            created.data["input_snapshot"]["request_intent"]["episode_ids"],
            sorted([str(first.id), str(second.id)]),
        )

    def test_confirmed_script_snapshot_detects_later_formal_drift(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="script-drift-run")
        run = ProductionRun.objects.get(pk=created.data["id"])
        source_stage = run.stages.get(stage_key="script_input")
        self.assertEqual(source_stage.status, "succeeded")
        self.assertEqual(len(source_stage.output_snapshot_sha256), 64)
        ScriptRevision.objects.filter(pk=revision_id).update(
            content="# 漂移",
            content_size_bytes=len("# 漂移".encode("utf-8")),
            content_sha256="f" * 64,
        )
        run.refresh_from_db()
        with self.assertRaises(StageBlocked) as error:
            _assert_frozen_formal_objects(run)
        self.assertEqual(error.exception.code, "script_snapshot_drift")

    def _freeze_storyboard_for_drift_test(self, *, key: str):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key=key)
        run = ProductionRun.objects.get(pk=created.data["id"])
        episode = Episode.objects.create(prompt_settings={"include_audio": True},
            project=self.project,
            number=1,
            title="冻结分镜",
            status=Episode.Status.DRAFT,
        )
        shot = Shot.objects.create(include_audio=True,
            episode=episode,
            number=1,
            title="第一镜",



            duration_ms=5000,
            media_type=Shot.MediaType.VIDEO,
        )
        canonical = _canonical_storyboard(run)
        stage = run.stages.get(stage_key="storyboard_review")
        assign_json_contract(stage, "output_snapshot", {
            "schema_version": 1,
            "stage_key": stage.stage_key,
            "output_reference": "storyboard-review:approved",
            "formal_object_type": "storyboard",
            "canonical": canonical,
            "canonical_sha256": _sha256(canonical),
        })
        stage.status = StageRun.Status.SUCCEEDED
        stage.save(update_fields=(
            "status", "output_snapshot", "output_snapshot_schema_version",
            "output_snapshot_size_bytes", "output_snapshot_sha256", "updated_at",
        ))
        return run, episode, shot

    def test_storyboard_lifecycle_status_change_is_not_formal_drift(self):
        run, episode, _shot = self._freeze_storyboard_for_drift_test(
            key="storyboard-status-projection",
        )

        episode.status = Episode.Status.COMPLETED
        episode.save(update_fields=("status", "updated_at"))

        _assert_frozen_formal_objects(run)

    def test_storyboard_creative_change_remains_formal_drift(self):
        run, _episode, shot = self._freeze_storyboard_for_drift_test(
            key="storyboard-creative-drift",
        )

        shot.title = "Shen Yan leaves the approved frame."
        shot.save(update_fields=("title", "updated_at"))

        with self.assertRaises(StageBlocked) as error:
            _assert_frozen_formal_objects(run)
        self.assertEqual(error.exception.code, "storyboard_snapshot_drift")

    def test_owner_project_config_api_rejects_secret_and_system_only_keys(self):
        project_key = ConfigDefinition(
            key="project.test.threshold",
            domain="production",
            value_type=ConfigDefinition.ValueType.INTEGER,
            allowed_scopes="system,project",
        )
        assign_json_contract(project_key, "value_schema", {"minimum": 1, "maximum": 9})
        project_key.save()
        secret_key = ConfigDefinition(
            key="project.test.secret",
            domain="provider",
            value_type=ConfigDefinition.ValueType.STRING,
            allowed_scopes="project",
            is_secret=True,
        )
        assign_json_contract(secret_key, "value_schema", {})
        secret_key.save()
        system_key = ConfigDefinition(
            key="project.test.system-only",
            domain="runtime",
            value_type=ConfigDefinition.ValueType.INTEGER,
            allowed_scopes="system",
        )
        assign_json_contract(system_key, "value_schema", {})
        system_key.save()

        published = self.client.put(f"/api/v1/projects/{self.project.id}/config/", {
            "values": {project_key.key: 4},
            "change_summary": "owner override",
        }, format="json")
        self.assertEqual(published.status_code, 200, published.data)
        self.assertEqual(published.data["effective"]["values"][project_key.key], 4)
        effective = published.data["effective"]
        self.assertNotIn("full_snapshot_sha256", effective)
        self.assertNotIn("revision_ids", effective)
        source = effective["sources"][project_key.key]
        self.assertFalse({
            "id", "source_id", "revision_id", "hash", "sha256",
        } & set(source))
        revision = self.project.config_revisions.get(status="published")
        self.assertEqual(revision.created_by_id, self.user.id)

        for rejected_key in (secret_key.key, system_key.key):
            rejected = self.client.patch(f"/api/v1/projects/{self.project.id}/config/", {
                "values": {rejected_key: 1},
            }, format="json")
            self.assertEqual(rejected.status_code, 400, rejected.data)
        foreign = get_user_model().objects.create_user(
            username="project-config-foreigner", password="StrongPass123!",
        )
        foreign_client = APIClient(); foreign_client.force_authenticate(foreign)
        self.assertEqual(
            foreign_client.get(f"/api/v1/projects/{self.project.id}/config/").status_code,
            404,
        )

    def test_owner_stage_projection_redacts_routing_config_and_nested_secrets(self):
        revision_id = self._confirmed_script()
        created = self._create_run(revision_id, key="public-stage-projection")
        self.assertEqual(created.status_code, 201, created.data)
        run = ProductionRun.objects.get(pk=created.data["id"])
        stage = run.stages.first()
        assign_json_contract(stage, "routing", {
            "schema_version": 1,
            "mode": "fixed",
            "routes": [{
                "index": 0,
                "workflow_version_id": "internal-workflow-version",
                "model_deployment_id": "internal-deployment",
                "provider_id": "internal-provider",
                "model_identifier": "public-model",
                "driver_type": "rest_async",
                "is_external": True,
                "effective_timeout_seconds": 60,
                "execution_config": {"api_key": "route-secret"},
                "credential_binding_sha256": "a" * 64,
            }],
        })
        assign_json_contract(stage, "routing_decision", {
            "schema_version": 1,
            "state": "selected",
            "selected_route_index": 0,
            "nested": {"authorization": "Bearer hidden"},
        })
        assign_json_contract(stage, "config_snapshot", {
            "schema_version": 1,
            "base": {
                "schema_version": 1,
                "values": {"provider.secret": {"credential_binding_sha256": "b" * 64}},
                "sources": {"provider.secret": {
                    "source": "system", "source_id": "secret-revision",
                    "revision_id": "secret-revision", "sha256": "c" * 64,
                }},
                "revision_ids": ["secret-revision"],
                "sha256": "d" * 64,
            },
            "routes": {},
        })
        assign_json_contract(stage, "input_snapshot", {
            "nested": {"token": "input-secret"},
        })
        assign_json_contract(stage, "output_snapshot", {
            "nested": {"password": "output-secret"},
        })
        stage.save()

        response = self.client.get(f"/api/v1/production-runs/{run.id}/")
        self.assertEqual(response.status_code, 200, response.data)
        projected = next(item for item in response.data["stages"] if item["id"] == str(stage.id))
        serialized = json.dumps(projected, ensure_ascii=False)
        for forbidden in (
            "route-secret", "input-secret", "output-secret", "secret-revision",
            "internal-workflow-version", "internal-deployment", "internal-provider",
            "credential_binding_sha256",
        ):
            self.assertNotIn(forbidden, serialized)
        self.assertEqual(projected["routing"]["routes"][0]["index"], 0)
        self.assertEqual(projected["routing"]["routes"][0]["model_identifier"], "public-model")

    def test_pdf_is_explicitly_rejected(self):
        upload = SimpleUploadedFile(
            "script.pdf",
            b"%PDF-1.7 not accepted",
            content_type="application/pdf",
        )

        response = self.client.post("/api/v1/script-sources/", {
            "project_id": str(self.project.id),
            "file": upload,
        }, format="multipart")

        self.assertEqual(response.status_code, 422, response.data)
        self.assertEqual(response.data["error"]["code"], "pdf_not_supported")
        self.assertFalse(ScriptSource.objects.exists())

    def test_docx_upload_extracts_preview_before_confirmation(self):
        from docx import Document

        buffer = BytesIO()
        document = Document()
        document.add_heading("第一集", level=1)
        document.add_paragraph("雨夜，沈砚推门而入。")
        document.save(buffer)
        upload = SimpleUploadedFile(
            "script.docx",
            buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

        response = self.client.post("/api/v1/script-sources/", {
            "project_id": str(self.project.id),
            "file": upload,
        }, format="multipart")

        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(response.data["source"]["source_format"], ScriptSource.SourceFormat.DOCX)
        self.assertIn("沈砚", response.data["revision"]["content"])
        self.assertEqual(response.data["revision"]["status"], ScriptRevision.Status.PREVIEW)
